import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows

# 示例数据框
df = pd.read_excel(r'D:\2019plan\周销售数据.xlsx')
df["周数"] = df["周数"].astype(int)

# 获取本周销量最大的10个SKU
current_week = df["周数"].max()
top_skus = df[df["周数"] == current_week].nlargest(10, "Units Ordered")["SKU"].tolist()

# 创建一个Excel工作簿
wb = Workbook()
filename = r'D:\\运营\\SKU_charts_filtered.xlsx'

# 输入一个整数，只取小于这个整数的周数的值
max_week_number = 12

# 为每个SKU创建一个Excel工作表并插入图表
for sku in top_skus:
    sku_df = df[(df['SKU'] == sku) & (df['周数'] < max_week_number)]

    # 创建一个新的工作表
    ws = wb.create_sheet(title=str(sku))

    # 将SKU数据框写入工作表
    for r in dataframe_to_rows(sku_df, index=False, header=True):
        ws.append(r)

    # 创建一个LineChart图表
    chart = LineChart()
    chart.title = f'SKU: {sku}'
    chart.x_axis.title = 'Week Number'
    chart.y_axis.title = 'Value'

    # 设置数据范围
    data = Reference(ws, min_col=9, min_row=1, max_col=9, max_row=len(sku_df)+1)
    units_ordered = Reference(ws, min_col=6, min_row=2, max_col=6, max_row=len(sku_df)+1)
    sessions_total = Reference(ws, min_col=2, min_row=2, max_col=2, max_row=len(sku_df)+1)

    # 添加数据系列到图表
    chart.add_data(units_ordered, titles_from_data=True)
    chart.add_data(sessions_total, titles_from_data=True)

    # 设置类别轴数据
    chart.set_categories(data)

    # 添加图表到工作表并设置图表位置
    ws.add_chart(chart, "K2")

# 保存Excel文件
wb.save(filename)

